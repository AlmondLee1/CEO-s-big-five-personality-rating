"""Collect independent API ratings of public CEO impressions."""
from __future__ import annotations

import argparse
import csv
import hashlib
import importlib
import json
import os
import re
import sys
import time
import tomllib
from datetime import datetime, timezone
from itertools import cycle
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent
DIMENSIONS = ("openness", "conscientiousness", "extraversion", "agreeableness", "neuroticism")
FAILURE_SCORE = 999.00
PROMPT_TEMPLATE = """Please rate the {{dimension}} personality trait of {{ceo_name}}, CEO of {{company}}.

Rating range: 1 to 7 (1 is the lowest; 7 is the highest).

Requirements:
1. Output exactly one numeric score with two decimal places.
2. Do not output any explanation, analysis, or additional text.
3. The score must be between 1 and 7, inclusive.

Score:"""


def digest(data):
    return hashlib.sha256(data).hexdigest()


def build_prompt(company, name, dimension):
    if dimension not in DIMENSIONS:
        raise ValueError(f"Unknown dimension: {dimension}")
    return (PROMPT_TEMPLATE.replace("{{company}}", company)
            .replace("{{ceo_name}}", name).replace("{{dimension}}", dimension))


def parse_score(text):
    value = text.strip()
    if re.fullmatch(r"(?:[1-6]\.[0-9]{2}|7\.00)", value):
        return float(value), "ok"
    return FAILURE_SCORE, "invalid_response"


def read_ceos(path):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    if not rows:
        raise ValueError("CEO input is empty")
    ids = set()
    for row in rows:
        if not all(row.get(field, "").strip() for field in ("ceo_id", "name", "company")):
            raise ValueError("Each CEO needs ceo_id, name, and company")
        if row["ceo_id"] in ids:
            raise ValueError(f"Duplicate CEO ID: {row['ceo_id']}")
        ids.add(row["ceo_id"])
    return rows


def load_config(path):
    with path.open("rb") as handle:
        config = tomllib.load(handle)
    run = config["run"]
    if any(int(run[key]) < 1 for key in ("iterations", "max_retries", "max_tokens")):
        raise ValueError("iterations, max_retries, and max_tokens must be positive")
    if float(run["request_delay_seconds"]) < 0 or float(run["timeout_seconds"]) <= 0:
        raise ValueError("Invalid delay or timeout")
    return config


def load_pilot_settings(path):
    with path.open("rb") as handle:
        pilot = tomllib.load(handle)["pilot"]
    for field in ("provider", "dimension", "input", "output_dir", "iterations"):
        if not str(pilot.get(field, "")).strip():
            raise ValueError(f"Set pilot.{field} in {path}")
    for field in ("provider", "dimension", "input", "output_dir"):
        if not isinstance(pilot[field], str):
            raise ValueError(f"pilot.{field} must be text in {path}")
    if int(pilot["iterations"]) < 1:
        raise ValueError("pilot.iterations must be positive")
    output_name = pilot["output_dir"]
    for field in ("provider", "dimension", "iterations"):
        output_name = output_name.replace("{" + field + "}", str(pilot[field]))
    if "{" in output_name or "}" in output_name:
        raise ValueError("pilot.output_dir supports only {provider}, {dimension}, and {iterations}")
    pilot["input"] = (path.parent / pilot["input"]).resolve()
    pilot["output_dir"] = (path.parent / output_name).resolve()
    return pilot


def load_local_keys(path, key_names):
    if not path.is_file():
        return
    for line in path.read_text(encoding="utf-8-sig").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        name, value = line.split("=", 1)
        name, value = name.strip(), value.strip()
        if value.startswith(("'", '"')) and value.endswith(value[0]):
            value = value[1:-1]
        if name in key_names and value:
            os.environ.setdefault(name, value)


def check_dependencies(provider):
    modules = ("openpyxl", "anthropic" if provider == "claude" else
               "google.genai" if provider == "gemini" else "openai")
    for module in modules:
        try:
            importlib.import_module(module)
        except ImportError as exc:
            requirement = "requirements-optional.txt" if provider in ("claude", "gemini") and module != "openpyxl" else "requirements.txt"
            raise ValueError(f"Missing {module}; run: python -m pip install -r {requirement}") from exc


def make_api_call(provider, settings, run):
    model = settings["model"].strip()
    if not model:
        raise ValueError(f"Set the {provider} model ID in config.toml")
    key_name = settings["api_key_env"]
    key = os.getenv(key_name, "")
    if not key or key.startswith("your_"):
        raise ValueError(f"Set {key_name} in the environment or local .env")
    timeout = float(run["timeout_seconds"])
    tokens = int(run["max_tokens"])
    temperature = run["temperature"]
    adapter = settings["adapter"]
    if adapter == "openai_compatible":
        from openai import OpenAI
        client = OpenAI(api_key=key, base_url=settings["base_url"], timeout=timeout, max_retries=0)

        def call(prompt):
            options = {"model": model, "messages": [{"role": "user", "content": prompt}],
                       ("max_completion_tokens" if provider == "gpt" else "max_tokens"): tokens}
            if temperature is not None:
                options["temperature"] = float(temperature)
            if provider == "deepseek":
                options["extra_body"] = {"thinking": {"type": run["thinking_mode"]}}
            result = client.chat.completions.create(**options)
            content = result.choices[0].message.content if result.choices else ""
            return content or "", result.model or model
    elif adapter == "anthropic":
        from anthropic import Anthropic
        client = Anthropic(api_key=key, timeout=timeout, max_retries=0)

        def call(prompt):
            options = {"model": model, "messages": [{"role": "user", "content": prompt}],
                       "max_tokens": tokens}
            if temperature is not None:
                options["temperature"] = float(temperature)
            result = client.messages.create(**options)
            return "".join(block.text for block in result.content if block.type == "text"), result.model
    elif adapter == "gemini":
        from google import genai
        from google.genai import types
        client = genai.Client(api_key=key, http_options=types.HttpOptions(timeout=int(timeout * 1000)))

        def call(prompt):
            result = client.models.generate_content(
                model=model, contents=prompt,
                config=types.GenerateContentConfig(
                    temperature=float(temperature) if temperature is not None else None,
                    max_output_tokens=tokens, candidate_count=1))
            return result.text or "", getattr(result, "model_version", None) or model
    else:
        raise ValueError(f"Unknown adapter: {adapter}")
    return call


def call_with_retries(call, prompt, max_retries, on_retry=None):
    for attempt in range(1, max_retries + 1):
        try:
            response, model = call(prompt)
            return response, model, attempt, None
        except Exception as exc:
            error = type(exc).__name__
            if attempt < max_retries:
                if on_retry is not None:
                    on_retry(attempt, max_retries, error)
                time.sleep(min(2 ** (attempt - 1), 30))
    return "", None, max_retries, error


def run_metadata(dimension, iterations, input_hash, provider, model, run, demo):
    return {
        "schema_version": 2, "provider": provider, "model_requested": model,
        "dimension": dimension, "iterations_per_ceo": iterations,
        "input_sha256": input_hash, "prompt_sha256": digest(PROMPT_TEMPLATE.encode()),
        "temperature_requested": run["temperature"], "max_tokens": int(run["max_tokens"]),
        "thinking_mode": run.get("thinking_mode", "provider_default"),
        "max_retries": int(run["max_retries"]), "timeout_seconds": float(run["timeout_seconds"]),
        "request_delay_seconds": float(run["request_delay_seconds"]),
        "mode": "synthetic_demo" if demo else "live_api",
    }


def check_existing_run(output_dir, metadata, resume):
    log = output_dir / "ratings.jsonl"
    manifest = output_dir / "run.json"
    if not (log.exists() or manifest.exists()):
        return
    if not resume or not log.exists() or not manifest.exists():
        raise FileExistsError("Output exists; choose a new output_dir or resume the unchanged run")
    old = json.loads(manifest.read_text(encoding="utf-8"))
    changed = [key for key, value in metadata.items() if old.get(key) != value]
    if changed:
        raise ValueError(
            f"Saved run has different settings ({', '.join(changed)}). "
            "Restore the original settings to resume it, or set a new output_dir in pilot.toml."
        )


def show_progress(recorded, total, valid):
    width = 20
    filled = width * recorded // total
    bar = "#" * filled + "-" * (width - filled)
    print(f"Progress: [{bar}] {recorded}/{total} ({recorded / total:.0%}); "
          f"valid {valid}; failed {recorded - valid}", flush=True)


def write_outputs(records, ceos, iterations, output_dir):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    slots = {(str(r["ceo_id"]), int(r["iteration"])): r for r in records}
    book = Workbook()
    sheet = book.active
    sheet.title = "Ratings"
    sheet.append(["Iteration"] + [f"{c['name']} [{c['ceo_id']}]" for c in ceos])
    for i in range(1, iterations + 1):
        sheet.append([i] + [slots.get((c["ceo_id"], i), {}).get("score") for c in ceos])
    means = []
    summary = book.create_sheet("Summary")
    summary.append(["CEO ID", "CEO", "Company", "Recorded", "Valid", "Failed", "Mean valid score"])
    csv_rows = []
    for c in ceos:
        rows = [slots[(c["ceo_id"], i)] for i in range(1, iterations + 1)
                if (c["ceo_id"], i) in slots]
        valid = [r["score"] for r in rows if r["status"] == "ok"]
        mean = round(sum(valid) / len(valid), 2) if valid else None
        means.append(mean)
        values = [c["ceo_id"], c["name"], c["company"], len(rows),
                  len(valid), len(rows) - len(valid), mean]
        summary.append(values)
        csv_rows.append(values)
    sheet.append(["Mean valid score"] + means)
    for row in sheet.iter_rows(min_row=2, min_col=2):
        for cell in row:
            cell.number_format = "0.00"
    for tab in (sheet, summary):
        for cell in tab[1]:
            cell.font = Font(bold=True)
    book.save(output_dir / "ratings.xlsx")
    with (output_dir / "summary.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["ceo_id", "ceo_name", "company", "recorded", "valid", "failed", "mean_valid_score"])
        writer.writerows(csv_rows)


def run_scores(ceos, dimension, iterations, output_dir, input_hash,
               provider, model, call, run, resume=False, demo=False):
    if iterations < 1:
        raise ValueError("iterations must be positive")
    output_dir.mkdir(parents=True, exist_ok=True)
    log = output_dir / "ratings.jsonl"
    manifest = output_dir / "run.json"
    metadata = run_metadata(dimension, iterations, input_hash, provider, model, run, demo)
    check_existing_run(output_dir, metadata, resume)
    if not log.exists():
        manifest.write_text(json.dumps({**metadata, "created_utc": datetime.now(timezone.utc).isoformat()},
                                       indent=2) + "\n", encoding="utf-8")
        log.touch()
    with log.open(encoding="utf-8") as handle:
        records = [json.loads(line) for line in handle if line.strip()]
    completed = {(str(r["ceo_id"]), int(r["iteration"])) for r in records}
    if len(completed) != len(records):
        raise ValueError("Duplicate CEO/iteration slot in saved records")
    total_slots = len(ceos) * iterations
    valid_count = sum(record["status"] == "ok" for record in records)
    show_progress(len(completed), total_slots, valid_count)
    new_slots = attempts_total = 0
    with log.open("a", encoding="utf-8", newline="\n") as handle:
        for ceo in ceos:
            for iteration in range(1, iterations + 1):
                slot = (ceo["ceo_id"], iteration)
                if slot in completed:
                    continue
                print(f"Requesting {len(completed) + 1}/{total_slots}: "
                      f"{ceo['name']}, repetition {iteration}/{iterations} "
                      f"(timeout {run['timeout_seconds']}s per attempt)...", flush=True)
                prompt = build_prompt(ceo["company"], ceo["name"], dimension)
                response, returned_model, attempts, error = call_with_retries(
                    call, prompt, int(run["max_retries"]),
                    on_retry=lambda attempt, maximum, kind: print(
                        f"  Attempt {attempt}/{maximum} failed ({kind}); retrying...", flush=True))
                score, status = parse_score(response) if error is None else (FAILURE_SCORE, "api_error")
                record = {
                    "ceo_id": ceo["ceo_id"], "ceo_name": ceo["name"],
                    "company": ceo["company"], "dimension": dimension,
                    "provider": provider, "model_requested": model,
                    "model_returned": returned_model, "iteration": iteration,
                    "score": score, "status": status,
                    "error_type": error or ("format_or_range" if status == "invalid_response" else None),
                    "api_attempts": 0 if demo else attempts,
                }
                handle.write(json.dumps(record, ensure_ascii=False) + "\n")
                handle.flush()
                records.append(record)
                completed.add(slot)
                valid_count += status == "ok"
                new_slots += 1
                attempts_total += 0 if demo else attempts
                show_progress(len(completed), total_slots, valid_count)
                if not demo and float(run["request_delay_seconds"]) > 0:
                    time.sleep(float(run["request_delay_seconds"]))
    write_outputs(records, ceos, iterations, output_dir)
    return new_slots, len(records), attempts_total


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=BASE_DIR / "config.toml")
    sub = parser.add_subparsers(dest="command", required=True)
    demo = sub.add_parser("demo", help="Generate a synthetic example without API calls")
    demo.add_argument("--input", type=Path, default=BASE_DIR / "data/ceo_pilot_10.csv")
    demo.add_argument("--iterations", type=int, default=2)
    demo.add_argument("--output-dir", type=Path, required=True)
    demo.add_argument("--resume", action="store_true")
    pilot = sub.add_parser("pilot", help="Run the editable pilot.toml settings")
    pilot.add_argument("--settings", type=Path, default=BASE_DIR / "pilot.toml")
    pilot.add_argument("--check", action="store_true", help="Validate settings without API calls")
    pilot.add_argument("--resume", action="store_true")
    pilot.add_argument("--confirm-live-api", action="store_true")
    live = sub.add_parser("score", help="Make independent live API calls")
    live.add_argument("--provider", required=True)
    live.add_argument("--dimension", choices=DIMENSIONS, required=True)
    live.add_argument("--input", type=Path, default=BASE_DIR / "data/ceo_pilot_10.csv")
    live.add_argument("--iterations", type=int)
    live.add_argument("--output-dir", type=Path, required=True)
    live.add_argument("--resume", action="store_true")
    live.add_argument("--confirm-live-api", action="store_true")
    args = parser.parse_args()
    config = load_config(args.config)
    run = config["run"]
    if args.command == "demo":
        input_path = args.input
        ceos = read_ceos(input_path)
        responses = cycle(("4.25", "I cannot rate this person.", "5.00", "6.00"))

        def call(_):
            return next(responses), "synthetic"

        result = run_scores(ceos, "extraversion", args.iterations, args.output_dir,
                            digest(input_path.read_bytes()), "synthetic", "synthetic",
                            call, run, resume=args.resume, demo=True)
    else:
        if args.command == "pilot":
            settings_file = args.settings.resolve()
            pilot_settings = load_pilot_settings(settings_file)
            provider = pilot_settings["provider"]
            dimension = pilot_settings["dimension"]
            input_path = pilot_settings["input"]
            iterations = int(pilot_settings["iterations"])
            output_dir = pilot_settings["output_dir"]
        else:
            provider = args.provider
            dimension = args.dimension
            input_path = args.input
            iterations = args.iterations if args.iterations is not None else int(run["iterations"])
            output_dir = args.output_dir
        if not args.confirm_live_api and not (args.command == "pilot" and args.check):
            parser.error("Live calls require --confirm-live-api")
        if provider not in config["providers"]:
            parser.error(f"Unknown provider: {provider}")
        if dimension not in DIMENSIONS:
            parser.error(f"Unknown dimension: {dimension}")
        if not input_path.is_file():
            parser.error(f"CEO input CSV not found: {input_path}")
        load_local_keys(BASE_DIR / ".env", {item["api_key_env"] for item in config["providers"].values()})
        settings = config["providers"][provider]
        effective_run = dict(run)
        effective_run["max_tokens"] = int(settings.get("max_tokens", run["max_tokens"]))
        if settings.get("temperature_mode") == "omit":
            effective_run["temperature"] = None
        effective_run["thinking_mode"] = settings.get("thinking_mode", "provider_default")
        ceos = read_ceos(input_path)
        if args.command == "pilot":
            if args.check:
                expected = run_metadata(dimension, iterations, digest(input_path.read_bytes()),
                                        provider, settings["model"], effective_run, False)
                check_existing_run(output_dir, expected, resume=True)
            check_dependencies(provider)
            key_name = settings["api_key_env"]
            key = os.getenv(key_name, "")
            print(f"Pilot: {len(ceos)} CEOs x {iterations} repetitions x 1 trait = {len(ceos) * iterations} rating slots", flush=True)
            print(f"Provider: {provider}; model: {settings['model']}; trait: {dimension}", flush=True)
            print(f"Input: {input_path}; output: {output_dir}", flush=True)
            print(f"API key variable: {key_name}; configured: {'yes' if key and not key.startswith('your_') else 'no'}", flush=True)
            if not key or key.startswith("your_"):
                raise ValueError(
                    f"{key_name} is missing. Copy .env.example to .env in this folder, "
                    f"replace the {key_name} placeholder with your own key, save, and rerun."
                )
            if args.check:
                return 0
        result = run_scores(ceos, dimension, iterations,
                            output_dir, digest(input_path.read_bytes()), provider,
                            settings["model"], make_api_call(provider, settings, effective_run),
                            effective_run, resume=args.resume)
    print(f"New slots: {result[0]}; total records: {result[1]}; API attempts: {result[2]}; output: {args.output_dir if args.command == 'demo' else output_dir}")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except (FileNotFoundError, FileExistsError, KeyError, ValueError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        sys.exit(2)
