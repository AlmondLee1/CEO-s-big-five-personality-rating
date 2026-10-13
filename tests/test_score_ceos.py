"""Check rating slots, refusal handling, spreadsheet output, and resume."""
import json
import io
import os
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook

from score_ceos import FAILURE_SCORE, build_prompt, main, make_api_call, parse_score, read_ceos, run_scores


RUN = {"temperature": 0.4, "max_tokens": 20, "max_retries": 2,
       "timeout_seconds": 30, "request_delay_seconds": 0}


class RatingTests(unittest.TestCase):
    def test_public_pilot_input_and_relative_paths(self):
        package = Path(__file__).resolve().parents[1]
        self.assertEqual(len(read_ceos(package / "data" / "ceo_pilot_10.csv")), 10)
        root = package / "outputs"
        root.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=root) as temporary:
            folder = Path(temporary)
            (folder / "input.csv").write_text(
                "ceo_id,name,company\n1,First CEO,First Company\n2,Second CEO,Second Company\n",
                encoding="utf-8")
            settings = folder / "pilot.toml"
            settings.write_text(
                '[pilot]\nprovider="deepseek"\ndimension="extraversion"\niterations=2\n'
                'input="input.csv"\noutput_dir="results-{provider}-{dimension}-{iterations}"\n',
                encoding="utf-8")
            argv = ["score_ceos.py", "pilot", "--settings", str(settings), "--confirm-live-api"]
            with patch.dict(os.environ, {"DEEPSEEK_API_KEY": "test-key"}):
                with patch("score_ceos.make_api_call", return_value=lambda _: ("4.50", "fake-model")):
                    with patch.object(sys, "argv", argv):
                        self.assertEqual(main(), 0)
            result_dir = folder / "results-deepseek-extraversion-2"
            self.assertEqual(len((result_dir / "ratings.jsonl").read_text().splitlines()), 4)
            self.assertTrue((result_dir / "ratings.xlsx").is_file())
            settings.write_text(settings.read_text(encoding="utf-8").replace("iterations=2", "iterations=5"),
                                encoding="utf-8")
            check_argv = ["score_ceos.py", "pilot", "--settings", str(settings), "--check"]
            with patch.dict(os.environ, {"DEEPSEEK_API_KEY": ""}):
                with patch("score_ceos.load_local_keys"):
                    with patch("score_ceos.make_api_call") as api:
                        with patch.object(sys, "argv", check_argv):
                            with self.assertRaisesRegex(ValueError, "DEEPSEEK_API_KEY is missing"):
                                main()
                        api.assert_not_called()
            self.assertFalse((folder / "results-deepseek-extraversion-5").exists())

    def test_reasoning_provider_request_parameters(self):
        fake_client = MagicMock()
        fake_client.chat.completions.create.return_value = SimpleNamespace(
            choices=[SimpleNamespace(message=SimpleNamespace(content="4.00"))], model="served")
        with patch.dict(os.environ, {"DEEPSEEK_API_KEY": "test-key", "OPENAI_API_KEY": "test-key"}):
            with patch("openai.OpenAI", return_value=fake_client):
                deepseek = make_api_call("deepseek", {
                    "adapter": "openai_compatible", "model": "deepseek-flash",
                    "api_key_env": "DEEPSEEK_API_KEY", "base_url": "https://api.deepseek.com"},
                    {**RUN, "temperature": None, "max_tokens": 1024, "thinking_mode": "enabled"})
                self.assertEqual(deepseek("Prompt"), ("4.00", "served"))
                deepseek_args = fake_client.chat.completions.create.call_args.kwargs
                self.assertEqual(deepseek_args["max_tokens"], 1024)
                self.assertEqual(deepseek_args["extra_body"], {"thinking": {"type": "enabled"}})
                self.assertNotIn("temperature", deepseek_args)
                gpt = make_api_call("gpt", {
                    "adapter": "openai_compatible", "model": "gpt-6-luna",
                    "api_key_env": "OPENAI_API_KEY", "base_url": "https://api.openai.com/v1"},
                    {**RUN, "temperature": None, "max_tokens": 1024})
                gpt("Prompt")
                gpt_args = fake_client.chat.completions.create.call_args.kwargs
                self.assertEqual(gpt_args["max_completion_tokens"], 1024)
                self.assertNotIn("temperature", gpt_args)

    def test_prompt_and_format(self):
        prompt = build_prompt("Example Inc.", "Example CEO", "extraversion")
        self.assertIn("extraversion personality trait", prompt)
        self.assertIn("exactly one numeric score with two decimal places", prompt)
        self.assertEqual(parse_score("4.25"), (4.25, "ok"))
        for response in ("5", "7.01", "999.00", "Score: 5.00", "I cannot rate this person."):
            self.assertEqual(parse_score(response), (FAILURE_SCORE, "invalid_response"))

    def test_all_slots_and_resume(self):
        ceos = [{"ceo_id": "EX|Person", "name": "Person", "company": "Example Inc."}]
        responses = iter(("4.00", "I cannot rate this person.", "5.00"))

        def call(_):
            return next(responses), "model-version"

        root = Path(__file__).resolve().parents[1] / "outputs"
        root.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=root) as temporary:
            output = Path(temporary)
            args = (ceos, "extraversion", 3, output, "input-hash", "provider", "model", call, RUN)
            progress = io.StringIO()
            with redirect_stdout(progress):
                self.assertEqual(run_scores(*args), (3, 3, 3))
                self.assertEqual(run_scores(*args, resume=True), (0, 3, 0))
            self.assertIn("Progress: [####################] 3/3 (100%); valid 2; failed 1", progress.getvalue())
            self.assertIn("Requesting 1/3: Person, repetition 1/3", progress.getvalue())
            with self.assertRaisesRegex(ValueError, "iterations_per_ceo.*new output_dir"):
                run_scores(ceos, "extraversion", 1, output, "input-hash", "provider",
                           "model", call, RUN, resume=True)
            records = [json.loads(line) for line in (output / "ratings.jsonl").read_text().splitlines()]
            self.assertEqual([row["score"] for row in records], [4.0, 999.0, 5.0])
            book = load_workbook(output / "ratings.xlsx", read_only=True)
            sheet = book["Ratings"]
            self.assertEqual([sheet.cell(i, 2).value for i in (2, 3, 4, 5)], [4, 999, 5, 4.5])
            book.close()

    def test_api_error_does_not_stop_next_slot(self):
        ceos = [{"ceo_id": "1", "name": "Person", "company": "Example"}]
        counter = iter((True, True, False))

        def call(_):
            if next(counter):
                raise RuntimeError("Unavailable")
            return "6.00", "model-version"

        root = Path(__file__).resolve().parents[1] / "outputs"
        root.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=root) as temporary:
            output = Path(temporary)
            run_scores(ceos, "openness", 2, output, "hash", "provider", "model", call, RUN)
            records = [json.loads(line) for line in (output / "ratings.jsonl").read_text().splitlines()]
            self.assertEqual([row["status"] for row in records], ["api_error", "ok"])
            self.assertEqual([row["score"] for row in records], [999.0, 6.0])


if __name__ == "__main__":
    unittest.main()
